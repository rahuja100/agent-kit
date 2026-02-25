"""Tools for Sales Enablement Agent.

Pure utility functions -- no LLM logic. Each function handles one concern:
  - Excel reading (scan)
  - Client profile persistence (save/load JSON)
  - Analysis persistence (save/load JSON)
  - PPTX generation
"""

import json
import logging
import re
from datetime import datetime, UTC
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger(__name__)

AGENT_DIR = Path(__file__).parent
INPUT_DIR = AGENT_DIR / "input"
DATA_DIR = AGENT_DIR / "data"
OUTPUT_DIR = AGENT_DIR / "output"
TEMPLATES_DIR = AGENT_DIR / "templates"

CLIENT_LIST_FILE = INPUT_DIR / "client list.xlsx"
CLIENTS_JSON = DATA_DIR / "clients.json"
PROFILES_JSON = DATA_DIR / "client_profiles.json"
TEMPLATE_FILE = TEMPLATES_DIR / "template.pptx"


def _ensure_dirs() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)


# ── Step 1: Scan ──────────────────────────────────────────────────

def _normalize_client_name(name: str) -> str:
    """Strip engagement suffixes like '- 1', '- 2'."""
    return re.sub(r"\s*-\s*\d+\s*$", "", name).strip()


def _find_header_row(ws: Any, target_headers: set[str], max_scan: int = 20) -> tuple[int, dict[str, int]]:
    """Scan for the row containing known column headers. Returns (row_num, col_map)."""
    for r in range(1, min(ws.max_row + 1, max_scan)):
        col_map: dict[str, int] = {}
        for c in range(1, min(ws.max_column + 1, 50)):
            val = ws.cell(r, c).value
            if val:
                header = str(val).strip().lower()
                if header in target_headers:
                    col_map[header] = c
        if col_map:
            return r, col_map
    return 1, {}


def scan_client_list(file_path: Path | None = None) -> dict[str, Any]:
    """Read client names + PE sponsors from Excel, merge with existing clients.json."""
    path = file_path or CLIENT_LIST_FILE
    if not path.exists():
        return {"error": f"File not found: {path}"}

    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)

        ws = None
        for sheet_name in wb.sheetnames:
            if "contract" in sheet_name.lower():
                ws = wb[sheet_name]
                break
        if ws is None:
            ws = wb.active
        if ws is None:
            return {"error": "No worksheet found"}

        target_headers = {
            "portfolio company", "company", "company_", "client", "client name",
            "company name", "name", "full name", "pe client", "rf name",
        }
        header_row, col_map = _find_header_row(ws, target_headers)

        name_col = 0
        for preferred in ("portfolio company", "company_", "full name", "company name", "company", "client", "name"):
            if preferred in col_map:
                name_col = col_map[preferred]
                break
        if name_col == 0:
            name_col = 1

        pe_col = col_map.get("pe client", 0)

        skip_values = {
            "portfolio company", "company", "client", "client name", "company name",
            "name", "company_", "full name", "existing customers", "new customers",
        }
        seen: set[str] = set()
        new_clients: list[dict[str, str]] = []
        for r in range(header_row + 1, ws.max_row + 1):
            val = ws.cell(r, name_col).value
            if val is None:
                continue
            raw = str(val).strip()
            if not raw or raw.lower() in skip_values:
                continue
            name = _normalize_client_name(raw)
            if not name or name.lower() in seen:
                continue
            seen.add(name.lower())

            pe_sponsor = ""
            if pe_col:
                pe_val = ws.cell(r, pe_col).value
                if pe_val:
                    pe_sponsor = str(pe_val).strip()

            new_clients.append({"name": name, "pe_sponsor": pe_sponsor})
        wb.close()

        # Merge with existing clients.json
        _ensure_dirs()
        existing: list[dict[str, str]] = []
        if CLIENTS_JSON.exists():
            with open(CLIENTS_JSON) as f:
                data = json.load(f)
                existing = data.get("clients", [])

        existing_lower = {c["name"].lower() if isinstance(c, dict) else c.lower() for c in existing}
        added = [c for c in new_clients if c["name"].lower() not in existing_lower]
        merged = existing + added

        result = {
            "clients": merged,
            "count": len(merged),
            "last_scanned": datetime.now(UTC).isoformat(),
        }
        with open(CLIENTS_JSON, "w") as f:
            json.dump(result, f, indent=2)

        logger.info(f"Scan complete: {len(merged)} total ({len(added)} new, {len(existing)} existing)")
        return {"total": len(merged), "new": len(added), "existing": len(existing), "file": str(CLIENTS_JSON)}

    except Exception as e:
        return {"error": f"Failed to read Excel: {e}"}


# ── Step 2: Research (tool functions for LLM) ────────────────────

async def load_clients_for_research() -> dict[str, Any]:
    """Load clients.json and determine which ones still need research."""
    if not CLIENTS_JSON.exists():
        return {"error": "No clients.json found. Run /scan first."}

    with open(CLIENTS_JSON) as f:
        data = json.load(f)
    all_clients: list[dict[str, str]] = data.get("clients", [])

    # Check which are already profiled
    profiled_names: set[str] = set()
    if PROFILES_JSON.exists():
        with open(PROFILES_JSON) as f:
            existing = json.load(f)
        profiled_names = {p["name"].lower() for p in existing}

    pending = [c for c in all_clients if c["name"].lower() not in profiled_names]

    return {
        "total_clients": len(all_clients),
        "already_profiled": len(profiled_names),
        "needs_research": len(pending),
        "pending_clients": pending,
    }


async def save_single_client_profile(profile_json: str) -> dict[str, Any]:
    """Save one client profile, appending to client_profiles.json. Returns progress info."""
    _ensure_dirs()

    try:
        profile = json.loads(profile_json)
    except json.JSONDecodeError as e:
        return {"error": f"Invalid JSON: {e}"}

    if not isinstance(profile, dict) or "name" not in profile:
        return {"error": "profile_json must be a JSON object with at least a 'name' field"}

    existing: list[dict[str, Any]] = []
    if PROFILES_JSON.exists():
        with open(PROFILES_JSON) as f:
            existing = json.load(f)

    # Skip if already profiled
    if any(p["name"].lower() == profile["name"].lower() for p in existing):
        return {"status": "skipped", "reason": "already profiled", "completed": len(existing), "name": profile["name"]}

    existing.append(profile)

    with open(PROFILES_JSON, "w") as f:
        json.dump(existing, f, indent=2)

    # Load total to calculate progress
    total = 0
    if CLIENTS_JSON.exists():
        with open(CLIENTS_JSON) as f:
            total = json.load(f).get("count", 0)

    completed = len(existing)
    remaining = max(0, total - completed)
    logger.info(f"Saved profile: {profile['name']} ({completed}/{total})")

    return {
        "status": "saved",
        "name": profile["name"],
        "completed": completed,
        "total": total,
        "remaining": remaining,
        "progress": f"{completed}/{total}",
    }


async def update_client_profile(name: str, updates_json: str) -> dict[str, Any]:
    """Update an existing client profile by name. Merges updates into the existing profile."""
    if not PROFILES_JSON.exists():
        return {"error": "No client_profiles.json found."}

    try:
        updates = json.loads(updates_json)
    except json.JSONDecodeError as e:
        return {"error": f"Invalid JSON: {e}"}

    with open(PROFILES_JSON) as f:
        profiles = json.load(f)

    found = False
    for p in profiles:
        if p["name"].lower() == name.lower():
            p.update(updates)
            found = True
            break

    if not found:
        return {"error": f"Client '{name}' not found in profiles"}

    with open(PROFILES_JSON, "w") as f:
        json.dump(profiles, f, indent=2)

    logger.info(f"Updated profile for {name}")
    return {"status": "updated", "name": name}


def replace_client_profile(old_name: str, new_profile: dict[str, Any]) -> dict[str, Any]:
    """Replace a profile by removing the old one and inserting the new one."""
    if not PROFILES_JSON.exists():
        return {"error": "No client_profiles.json found."}

    with open(PROFILES_JSON) as f:
        profiles = json.load(f)

    # Remove the old profile
    profiles = [p for p in profiles if p["name"].lower() != old_name.lower()]
    # Remove any existing profile with the new name too (avoid duplicates)
    new_name = new_profile.get("name", "")
    if new_name:
        profiles = [p for p in profiles if p["name"].lower() != new_name.lower()]

    profiles.append(new_profile)

    with open(PROFILES_JSON, "w") as f:
        json.dump(profiles, f, indent=2)

    logger.info(f"Replaced profile: '{old_name}' -> '{new_name}'")
    return {"status": "replaced", "old_name": old_name, "new_name": new_name}


def get_flagged_profiles() -> list[dict[str, Any]]:
    """Return profiles with low confidence only."""
    if not PROFILES_JSON.exists():
        return []
    with open(PROFILES_JSON) as f:
        profiles = json.load(f)
    return [p for p in profiles if p.get("confidence", "high") == "low"]


# ── Savings benchmarks ────────────────────────────────────────────

SAVINGS_BENCHMARKS_FILE = DATA_DIR / "savings_benchmarks.yaml"
REVENUE_BENCHMARKS_FILE = DATA_DIR / "revenue_benchmarks.yaml"


def load_savings_benchmarks() -> str:
    """Load savings benchmarks YAML as a string for injection into the prompt."""
    if not SAVINGS_BENCHMARKS_FILE.exists():
        return ""
    return SAVINGS_BENCHMARKS_FILE.read_text()


def load_revenue_benchmarks() -> str:
    """Load revenue estimation benchmarks YAML for injection into the prompt."""
    if not REVENUE_BENCHMARKS_FILE.exists():
        return ""
    return REVENUE_BENCHMARKS_FILE.read_text()


# ── Step 3: Analyze (tool functions for LLM) ─────────────────────

def _pe_slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def _working_path(pe_firm: str) -> Path:
    """Path to the in-progress working analysis file."""
    return DATA_DIR / f"working_{_pe_slug(pe_firm)}.json"


def load_client_profiles_for_comps() -> str:
    """Load client profiles as a compact string for injection into company analysis prompts."""
    if not PROFILES_JSON.exists():
        return "No client profiles found. Run /research first."

    with open(PROFILES_JSON) as f:
        profiles = json.load(f)

    lines = []
    for p in profiles:
        lines.append(f"- {p.get('name', '')} | {p.get('industry', '')} | {p.get('sector', '')} | {p.get('description', '')[:80]}")
    return "\n".join(lines)


async def save_portfolio_companies(companies_json: str, pe_firm_name: str) -> dict[str, Any]:
    """Save discovered portfolio companies to the working analysis file."""
    _ensure_dirs()

    try:
        companies = json.loads(companies_json)
    except json.JSONDecodeError as e:
        return {"error": f"Invalid JSON: {e}"}

    working = {
        "pe_firm": {"name": pe_firm_name},
        "portfolio_companies": companies,
        "company_analyses": [],
    }

    path = _working_path(pe_firm_name)
    with open(path, "w") as f:
        json.dump(working, f, indent=2)

    logger.info(f"Saved {len(companies)} portfolio companies to {path}")
    return {"status": "saved", "count": len(companies), "file": str(path)}


async def save_company_analysis(analysis_json: str, pe_firm_name: str) -> dict[str, Any]:
    """Save a single company's analysis (comps + savings) to the working file."""
    _ensure_dirs()

    try:
        analysis = json.loads(analysis_json)
    except json.JSONDecodeError as e:
        return {"error": f"Invalid JSON: {e}"}

    path = _find_working_file(pe_firm_name)
    if path is None:
        return {"error": "No working analysis file. Run portfolio search first."}

    with open(path) as f:
        working = json.load(f)

    # Append or replace this company's analysis
    existing = working.get("company_analyses", [])
    company_name = analysis.get("company_name", "")
    existing = [a for a in existing if a.get("company_name", "").lower() != company_name.lower()]
    existing.append(analysis)
    working["company_analyses"] = existing

    with open(path, "w") as f:
        json.dump(working, f, indent=2)

    total = len(working.get("portfolio_companies", []))
    completed = len(existing)
    logger.info(f"Saved company analysis: {company_name} ({completed}/{total})")
    return {
        "status": "saved",
        "company": company_name,
        "completed": completed,
        "total": total,
        "progress": f"{completed}/{total}",
    }


async def load_working_analysis(pe_firm_name: str) -> dict[str, Any]:
    """Load the working analysis file for summary generation."""
    path = _working_path(pe_firm_name)
    if not path.exists():
        return {"error": f"No working analysis for {pe_firm_name}."}

    with open(path) as f:
        return json.load(f)


async def finalize_analysis(pe_firm_name: str, pe_firm_json: str, executive_summary: str) -> dict[str, Any]:
    """Assemble the final analysis JSON from the working file and save to analysis_{slug}.json."""
    _ensure_dirs()
    path = _working_path(pe_firm_name)
    if not path.exists():
        return {"error": "No working analysis file found."}

    with open(path) as f:
        working = json.load(f)

    try:
        pe_firm_data = json.loads(pe_firm_json)
    except json.JSONDecodeError:
        pe_firm_data = working.get("pe_firm", {"name": pe_firm_name})

    # Load client industry mappings from profiles
    client_mappings = []
    if PROFILES_JSON.exists():
        with open(PROFILES_JSON) as f:
            profiles = json.load(f)
        client_mappings = [{"client_name": p.get("name", ""), "industry": p.get("industry", "")} for p in profiles]

    # Assemble from company analyses
    companies = working.get("company_analyses", [])
    final = {
        "pe_firm": pe_firm_data,
        "portfolio_companies": [
            {
                "name": c.get("company_name", ""),
                "industry": c.get("industry", ""),
                "description": c.get("description", ""),
                "estimated_revenue_mm": c.get("estimated_revenue_mm", 0),
                "revenue_source": c.get("revenue_source", "heuristic"),
            }
            for c in companies
        ],
        "client_industry_mappings": client_mappings,
        "comps_analysis": [
            {
                "portfolio_company": c.get("company_name", ""),
                "matching_clients": c.get("matching_clients", []),
            }
            for c in companies
        ],
        "savings_assessment": [
            {
                "portfolio_company": c.get("company_name", ""),
                "estimated_revenue_mm": c.get("estimated_revenue_mm", 0),
                "spend_categories": c.get("spend_categories", []),
                "total_estimated_savings_mm": c.get("total_estimated_savings_mm", 0),
                "key_opportunities": c.get("key_opportunities", ""),
            }
            for c in companies
        ],
        "executive_summary": executive_summary,
    }

    slug = _pe_slug(pe_firm_name)
    out_path = DATA_DIR / f"analysis_{slug}.json"
    with open(out_path, "w") as f:
        json.dump(final, f, indent=2)

    from rich.console import Console
    Console().print(f"\n  [bold green]Analysis saved to {out_path}[/bold green]\n")
    logger.info(f"Finalized analysis to {out_path}")
    return {"file": str(out_path), "status": "success"}


def _find_working_file(pe_firm: str) -> Path | None:
    """Find the working file, with fuzzy matching if exact slug doesn't match."""
    exact = _working_path(pe_firm)
    if exact.exists():
        return exact
    # Fuzzy: find any working file whose slug contains our slug or vice versa
    slug = _pe_slug(pe_firm)
    for f in DATA_DIR.glob("working_*.json"):
        file_slug = f.stem.replace("working_", "")
        if slug in file_slug or file_slug in slug:
            return f
    return None


def get_pending_portfolio_companies(pe_firm: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Return (all_companies, already_analyzed_names) from the working file."""
    path = _find_working_file(pe_firm)
    if path is None:
        return [], []
    with open(path) as f:
        working = json.load(f)
    all_cos = working.get("portfolio_companies", [])
    done = [a.get("company_name", "").lower() for a in working.get("company_analyses", [])]
    return all_cos, done


# ── Step 4: Report ────────────────────────────────────────────────

def load_analysis(pe_firm_name: str) -> dict[str, Any] | None:
    """Load analysis JSON for a PE firm. Returns None if not found."""
    slug = re.sub(r"[^a-z0-9]+", "_", pe_firm_name.lower()).strip("_")
    path = DATA_DIR / f"analysis_{slug}.json"
    if not path.exists():
        return None
    with open(path) as f:
        return json.load(f)


def generate_report_pptx(pe_firm_name: str) -> dict[str, Any]:
    """Generate PPTX from saved analysis. Returns output path or error."""
    from .pptx_generator import generate_report

    analysis = load_analysis(pe_firm_name)
    if analysis is None:
        slug = re.sub(r"[^a-z0-9]+", "_", pe_firm_name.lower()).strip("_")
        return {"error": f"No analysis found at data/analysis_{slug}.json. Run /analyze first."}

    if not TEMPLATE_FILE.exists():
        return {"error": f"Template not found: {TEMPLATE_FILE}. Drop your .pptx template in templates/"}

    _ensure_dirs()
    slug = re.sub(r"[^a-z0-9]+", "_", pe_firm_name.lower()).strip("_")
    out_path = OUTPUT_DIR / f"{slug}_report.pptx"

    try:
        slide_count = generate_report(analysis, str(TEMPLATE_FILE), str(out_path))
        logger.info(f"Generated {slide_count} slides at {out_path}")
        return {"output_path": str(out_path), "slide_count": slide_count, "status": "success"}
    except Exception as e:
        logger.error(f"PPTX generation failed: {e}", exc_info=True)
        return {"error": f"PPTX generation failed: {e}"}


# ── Portco-level analysis ─────────────────────────────────────────

async def save_portco_analysis(analysis_json: str, company_name: str) -> dict[str, Any]:
    """Save a single portco analysis to data/portco_{slug}.json."""
    _ensure_dirs()

    try:
        data = json.loads(analysis_json)
    except json.JSONDecodeError as e:
        return {"error": f"Invalid JSON: {e}"}

    slug = _pe_slug(company_name)
    out_path = DATA_DIR / f"portco_{slug}.json"
    with open(out_path, "w") as f:
        json.dump(data, f, indent=2)

    from rich.console import Console
    Console().print(f"\n  [bold green]Portco analysis saved to {out_path}[/bold green]\n")
    logger.info(f"Saved portco analysis to {out_path}")
    return {"file": str(out_path), "status": "success"}


def load_portco_analysis(company_name: str) -> dict[str, Any] | None:
    """Load portco analysis JSON. Returns None if not found."""
    slug = _pe_slug(company_name)
    path = DATA_DIR / f"portco_{slug}.json"
    if not path.exists():
        return None
    with open(path) as f:
        return json.load(f)


def generate_portco_report_pptx(company_name: str) -> dict[str, Any]:
    """Generate portco-level PPTX from saved analysis."""
    from .pptx_generator import generate_portco_report

    analysis = load_portco_analysis(company_name)
    if analysis is None:
        slug = _pe_slug(company_name)
        return {"error": f"No portco analysis found at data/portco_{slug}.json. Run /analyze-portco first."}

    if not TEMPLATE_FILE.exists():
        return {"error": f"Template not found: {TEMPLATE_FILE}. Drop your .pptx template in templates/"}

    _ensure_dirs()
    slug = _pe_slug(company_name)
    out_path = OUTPUT_DIR / f"portco_{slug}_report.pptx"

    try:
        slide_count = generate_portco_report(analysis, str(TEMPLATE_FILE), str(out_path))
        logger.info(f"Generated portco report: {slide_count} slides at {out_path}")
        return {"output_path": str(out_path), "slide_count": slide_count, "status": "success"}
    except Exception as e:
        logger.error(f"Portco PPTX generation failed: {e}", exc_info=True)
        return {"error": f"PPTX generation failed: {e}"}


def get_portco_analyzer_tool_definitions() -> list[dict[str, Any]]:
    """Tools for analyzing a single portfolio company standalone."""
    return [
        {"type": "web_search", "search_context_size": "medium"},
        {
            "type": "function",
            "name": "save_portco_analysis",
            "description": "Save this company's completed analysis. Call after all analysis (revenue, comps, savings) is complete.",
            "parameters": {
                "type": "object",
                "properties": {
                    "analysis_json": {
                        "type": "string",
                        "description": (
                            "JSON object with: company_name, industry, description, estimated_revenue_mm, "
                            "revenue_source, employee_count, matching_clients, spend_categories, "
                            "total_estimated_savings_mm, key_opportunities"
                        ),
                    },
                    "company_name": {"type": "string"},
                },
                "required": ["analysis_json", "company_name"],
            },
        },
    ]


# ── Tool definitions for LLM steps ───────────────────────────────

def get_researcher_tool_definitions() -> list[dict[str, Any]]:
    """Tools available during /research (Step 2)."""
    return [
        {"type": "web_search", "search_context_size": "high"},
        {
            "type": "function",
            "name": "load_clients_for_research",
            "description": "Load the client list with PE sponsor context. Returns pending_clients (need research) with name and pe_sponsor for each.",
            "parameters": {"type": "object", "properties": {}, "required": []},
        },
        {
            "type": "function",
            "name": "save_single_client_profile",
            "description": (
                "Save ONE client profile. Call this for EACH client after researching it. "
                "Returns progress (e.g., '15/111'). The profile is appended to the profiles file."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "profile_json": {
                        "type": "string",
                        "description": (
                            "JSON object with: name, pe_sponsor, industry, sector, description "
                            "(2-3 sentences), estimated_revenue_range, confidence (high/medium/low), "
                            "notes (optional, for flags/uncertainty)"
                        ),
                    },
                },
                "required": ["profile_json"],
            },
        },
    ]


def get_portfolio_finder_tool_definitions() -> list[dict[str, Any]]:
    """Tools for finding PE firm portfolio companies."""
    return [
        {"type": "web_search", "search_context_size": "medium"},
        {
            "type": "function",
            "name": "save_portfolio_companies",
            "description": "Save the list of active portfolio companies found. Call after web search is complete.",
            "parameters": {
                "type": "object",
                "properties": {
                    "companies_json": {
                        "type": "string",
                        "description": "JSON array of objects with: name, industry, description",
                    },
                    "pe_firm_name": {"type": "string"},
                },
                "required": ["companies_json", "pe_firm_name"],
            },
        },
    ]


def get_company_analyzer_tool_definitions() -> list[dict[str, Any]]:
    """Tools for analyzing a single portfolio company."""
    return [
        {"type": "web_search", "search_context_size": "medium"},
        {
            "type": "function",
            "name": "save_company_analysis",
            "description": "Save this company's analysis (revenue, comps, savings). Call after all analysis is complete for this company.",
            "parameters": {
                "type": "object",
                "properties": {
                    "analysis_json": {
                        "type": "string",
                        "description": (
                            "JSON object with: company_name, industry, description, estimated_revenue_mm, "
                            "revenue_source, matching_clients (array of client_name/similarity/rationale), "
                            "spend_categories (array of category/estimated_spend_mm/savings_opportunity_pct/estimated_savings_mm), "
                            "total_estimated_savings_mm, key_opportunities"
                        ),
                    },
                    "pe_firm_name": {"type": "string"},
                },
                "required": ["analysis_json", "pe_firm_name"],
            },
        },
    ]


def get_summary_tool_definitions() -> list[dict[str, Any]]:
    """Tools for generating the executive summary."""
    return [
        {
            "type": "function",
            "name": "load_working_analysis",
            "description": "Load all completed company analyses for the PE firm.",
            "parameters": {
                "type": "object",
                "properties": {"pe_firm_name": {"type": "string"}},
                "required": ["pe_firm_name"],
            },
        },
        {
            "type": "function",
            "name": "finalize_analysis",
            "description": "Save the final analysis with executive summary. Call after writing the summary.",
            "parameters": {
                "type": "object",
                "properties": {
                    "pe_firm_name": {"type": "string"},
                    "pe_firm_json": {"type": "string", "description": "JSON with: name, description, website"},
                    "executive_summary": {"type": "string", "description": "The executive summary text"},
                },
                "required": ["pe_firm_name", "pe_firm_json", "executive_summary"],
            },
        },
    ]


async def execute_tool(tool_name: str, tool_args: dict[str, Any]) -> dict[str, Any]:
    """Route tool calls to the appropriate function."""
    if tool_name == "load_clients_for_research":
        return await load_clients_for_research()
    if tool_name == "save_single_client_profile":
        return await save_single_client_profile(tool_args["profile_json"])
    if tool_name == "update_client_profile":
        return await update_client_profile(tool_args["name"], tool_args["updates_json"])
    if tool_name == "save_portfolio_companies":
        return await save_portfolio_companies(tool_args["companies_json"], tool_args["pe_firm_name"])
    if tool_name == "save_portco_analysis":
        return await save_portco_analysis(tool_args["analysis_json"], tool_args["company_name"])
    if tool_name == "save_company_analysis":
        return await save_company_analysis(tool_args["analysis_json"], tool_args["pe_firm_name"])
    if tool_name == "load_working_analysis":
        return await load_working_analysis(tool_args["pe_firm_name"])
    if tool_name == "finalize_analysis":
        return await finalize_analysis(tool_args["pe_firm_name"], tool_args["pe_firm_json"], tool_args["executive_summary"])
    raise ValueError(f"Unknown tool: {tool_name}")
